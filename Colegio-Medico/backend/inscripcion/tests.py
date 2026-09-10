from django.test import TestCase
from django.contrib.auth.models import User
from django.utils import timezone
from datetime import timedelta
import pptx
import io

from jornada.models import Jornada
from inscripcion.models import Inscripcion
from usuario.models import PerfilUsuario
from rest_framework.test import APIClient

class CertificadoPrefijoTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.jornada = Jornada.objects.create(
            nombre="Jornada Médica Anual",
            costo=100.0,
            fecha_inicio=timezone.now().date() - timedelta(days=5),
            fecha_final=timezone.now().date() - timedelta(days=1),
            uc=5,
            estado_jornada='A'
        )

    def test_certificado_doctor_masculino(self):
        user_doc_m = User.objects.create_user(
            username='dr_juan',
            first_name='Juan',
            last_name='Pérez',
            email='juan@example.com'
        )
        PerfilUsuario.objects.create(user=user_doc_m, tipo='D', genero='M')

        inscripcion = Inscripcion.objects.create(
            usuario=user_doc_m,
            jornada=self.jornada
        )

        self.client.force_authenticate(user=user_doc_m)
        response = self.client.get(f'/api/inscripcion/{inscripcion.id}/certificado/')
        self.assertEqual(response.status_code, 200)

        # Leer presentación generada
        prs = pptx.Presentation(io.BytesIO(response.content))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('Dr. JUAN PÉREZ' in t for t in textos))

    def test_certificado_doctor_femenino(self):
        user_doc_f = User.objects.create_user(
            username='dra_maria',
            first_name='María',
            last_name='González',
            email='maria@example.com'
        )
        PerfilUsuario.objects.create(user=user_doc_f, tipo='D', genero='F')

        inscripcion = Inscripcion.objects.create(
            usuario=user_doc_f,
            jornada=self.jornada
        )

        self.client.force_authenticate(user=user_doc_f)
        response = self.client.get(f'/api/inscripcion/{inscripcion.id}/certificado/')
        self.assertEqual(response.status_code, 200)

        prs = pptx.Presentation(io.BytesIO(response.content))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('Dra. MARÍA GONZÁLEZ' in t for t in textos))

    def test_certificado_estudiante(self):
        user_est = User.objects.create_user(
            username='est_carlos',
            first_name='Carlos',
            last_name='Rodríguez',
            email='carlos@example.com'
        )
        PerfilUsuario.objects.create(user=user_est, tipo='E', genero='M')

        inscripcion = Inscripcion.objects.create(
            usuario=user_est,
            jornada=self.jornada
        )

        self.client.force_authenticate(user=user_est)
        response = self.client.get(f'/api/inscripcion/{inscripcion.id}/certificado/')
        self.assertEqual(response.status_code, 200)

        prs = pptx.Presentation(io.BytesIO(response.content))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('CARLOS RODRÍGUEZ' in t for t in textos))
        self.assertFalse(any('Dr. CARLOS RODRÍGUEZ' in t or 'Dra. CARLOS RODRÍGUEZ' in t for t in textos))

    def test_certificado_estudiante_con_prefijo_manual_lo_remueve(self):
        user_est = User.objects.create_user(
            username='est_dr_ana',
            first_name='Dr. Ana',
            last_name='Pérez',
            email='ana_est@example.com'
        )
        PerfilUsuario.objects.create(user=user_est, tipo='E', genero='F')

        inscripcion = Inscripcion.objects.create(
            usuario=user_est,
            jornada=self.jornada
        )

        self.client.force_authenticate(user=user_est)
        response = self.client.get(f'/api/inscripcion/{inscripcion.id}/certificado/')
        self.assertEqual(response.status_code, 200)

        prs = pptx.Presentation(io.BytesIO(response.content))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('ANA PÉREZ' in t for t in textos))
        self.assertFalse(any('Dr.' in t or 'Dra.' in t for t in textos if 'ANA PÉREZ' in t))

    def test_certificado_doctor_con_prefijo_manual_no_duplica(self):
        user_doc = User.objects.create_user(
            username='dr_pedro_manual',
            first_name='Dr. Pedro',
            last_name='Gómez',
            email='pedro_manual@example.com'
        )
        PerfilUsuario.objects.create(user=user_doc, tipo='D', genero='M')

        inscripcion = Inscripcion.objects.create(
            usuario=user_doc,
            jornada=self.jornada
        )

        self.client.force_authenticate(user=user_doc)
        response = self.client.get(f'/api/inscripcion/{inscripcion.id}/certificado/')
        self.assertEqual(response.status_code, 200)

        prs = pptx.Presentation(io.BytesIO(response.content))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('Dr. PEDRO GÓMEZ' in t for t in textos))
        self.assertFalse(any('Dr. DR.' in t or 'Dr. Dr.' in t for t in textos))


class EnvioAutomaticoCertificadosTests(TestCase):
    """Pruebas del envío automático del certificado por correo al culminar la jornada."""

    def setUp(self):
        self.jornada = Jornada.objects.create(
            nombre="Jornada Científica de Cardiología",
            costo=100.0,
            fecha_inicio=timezone.now().date() - timedelta(days=5),
            fecha_final=timezone.now().date() - timedelta(days=1),  # ya culminó
            uc=5,
            estado_jornada='A'
        )
        self.usuario = User.objects.create_user(
            username='participante1',
            first_name='Luis',
            last_name='Ramírez',
            email='luis@example.com'
        )
        PerfilUsuario.objects.create(user=self.usuario, tipo='D', genero='M')
        self.inscripcion = Inscripcion.objects.create(
            usuario=self.usuario,
            jornada=self.jornada
        )

    def test_correo_con_certificado_adjunto(self):
        from django.core import mail
        from inscripcion.utils import enviar_certificado_email

        resultado = enviar_certificado_email(self.inscripcion)

        self.assertTrue(resultado)
        # Ignorar los correos de registro/inscripción enviados por señales
        correos_cert = [c for c in mail.outbox if 'certificado' in c.subject.lower()]
        self.assertEqual(len(correos_cert), 1)

        correo = correos_cert[0]
        self.assertEqual(correo.to, ['luis@example.com'])
        self.assertIn('certificado', correo.subject.lower())
        self.assertIn('certificado de participación', correo.body.lower())

        # El certificado .pptx viene adjunto
        self.assertEqual(len(correo.attachments), 1)
        nombre_adjunto, contenido, mimetype = correo.attachments[0]
        self.assertTrue(nombre_adjunto.startswith('Certificado_Luis_Ramírez'))
        self.assertEqual(
            mimetype,
            'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )
        # El adjunto es un PPTX válido
        prs = pptx.Presentation(io.BytesIO(contenido))
        textos = [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        self.assertTrue(any('Dr. LUIS RAMÍREZ' in t for t in textos))

    def test_comando_envia_solo_jornadas_culminadas(self):
        from django.core import mail
        from io import StringIO
        from django.core.management import call_command

        # Jornada que aún NO culmina: no debe recibir certificado
        jornada_activa = Jornada.objects.create(
            nombre="Jornada Futura",
            costo=50.0,
            fecha_inicio=timezone.now().date() + timedelta(days=1),
            fecha_final=timezone.now().date() + timedelta(days=3),
            uc=2,
            estado_jornada='A'
        )
        otro_usuario = User.objects.create_user(
            username='participante2',
            first_name='Ana',
            last_name='Torres',
            email='ana_torres@example.com'
        )
        PerfilUsuario.objects.create(user=otro_usuario, tipo='D', genero='F')
        Inscripcion.objects.create(usuario=otro_usuario, jornada=jornada_activa)

        out = StringIO()
        call_command('enviar_certificados', stdout=out)

        correos_cert = [c for c in mail.outbox if 'certificado' in c.subject.lower()]
        self.assertEqual(len(correos_cert), 1)  # solo la jornada culminada
        self.assertEqual(correos_cert[0].to, ['luis@example.com'])
        self.assertIn('1 certificado(s) pendientes', out.getvalue())

        # La inscripción quedó marcada como enviada
        self.inscripcion.refresh_from_db()
        self.assertTrue(self.inscripcion.certificado_enviado)

        # Ejecutar de nuevo no reenvía (no duplica correos)
        call_command('enviar_certificados', stdout=StringIO())
        correos_cert = [c for c in mail.outbox if 'certificado' in c.subject.lower()]
        self.assertEqual(len(correos_cert), 1)

    def test_comando_sin_pendientes(self):
        from io import StringIO
        from django.core.management import call_command

        # Sin jornada culminada no hay nada que enviar
        self.jornada.fecha_final = timezone.now().date() + timedelta(days=10)
        self.jornada.save()

        out = StringIO()
        call_command('enviar_certificados', stdout=out)
        self.assertIn('No hay certificados pendientes', out.getvalue())

    def test_jornada_que_culmina_hoy_envia_certificado(self):
        """El certificado se envía el mismo día en que culmina la jornada."""
        from django.core import mail
        from io import StringIO
        from django.core.management import call_command

        # Jornada que culmina HOY
        self.jornada.fecha_final = timezone.now().date()
        self.jornada.save()

        out = StringIO()
        call_command('enviar_certificados', stdout=out)

        correos_cert = [c for c in mail.outbox if 'certificado' in c.subject.lower()]
        self.assertEqual(len(correos_cert), 1)
        self.assertEqual(correos_cert[0].to, ['luis@example.com'])
        self.inscripcion.refresh_from_db()
        self.assertTrue(self.inscripcion.certificado_enviado)
