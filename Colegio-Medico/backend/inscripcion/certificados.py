import os
import re
from io import BytesIO

from django.conf import settings
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image as PILImage
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from .models import Inscripcion


def _construir_nombre_titulado(usuario):
    """Aplica el protocolo del gremio (Dr./Dra.) y roles sobre el nombre del usuario."""
    first_name = (usuario.first_name or "").strip()
    last_name = (usuario.last_name or "").strip()
    nombre_base = f"{first_name} {last_name}".strip() or usuario.username.strip()

    # Limpiar cualquier prefijo ingresado manualmente en el nombre
    nombre_limpio = re.sub(r'^(dr\.|dra\.|dr|dra)\s+', '', nombre_base, flags=re.IGNORECASE).strip()

    prefijo = ""
    try:
        if usuario.perfil.tipo == 'D':
            prefijo = "Dra. " if usuario.perfil.genero == 'F' else "Dr. "
    except Exception:
        pass

    nombre_titulado = f"{prefijo}{nombre_limpio.upper()}"

    # Rol de Coordinador(a) — se asigna únicamente desde el admin de Django.
    rol_certificado = ""
    try:
        if usuario.perfil.coordinador:
            rol_certificado = "Coordinadora" if usuario.perfil.genero == 'F' else "Coordinador"
    except Exception:
        pass

    if rol_certificado:
        nombre_titulado = f"{nombre_titulado} — {rol_certificado}"

    return nombre_titulado


def generar_certificado_pptx(inscripcion):
    """
    Genera el certificado (archivo .pptx en memoria) para una inscripción.
    Devuelve un BytesIO listo para descargarse o adjuntarse a un correo.
    Lanza ValueError si no existe la plantilla.
    """
    usuario = inscripcion.usuario
    jornada = inscripcion.jornada

    nombre_titulado = _construir_nombre_titulado(usuario)

    # Sociedad científica del participante
    sociedad_certificado = ""
    try:
        sociedad_certificado = (usuario.perfil.sociedad or "").strip()
    except Exception:
        pass

    fecha_inicio_formateada = jornada.fecha_inicio.strftime('%d/%m/%Y')
    fecha_final_formateada = jornada.fecha_final.strftime('%d/%m/%Y')

    ruta_plantilla = os.path.join(settings.BASE_DIR, 'static', 'templates', 'plantilla_certificado.pptx')
    if not os.path.exists(ruta_plantilla):
        raise ValueError(f"No se encontró la plantilla en {ruta_plantilla}")

    prs = Presentation(ruta_plantilla)

    # La nueva plantilla es una imagen de fondo a pantalla completa, sin campos
    # de texto. Insertamos el nombre del participante como un cuadro de texto
    # sobre la zona vacía central de la imagen.
    for slide in prs.slides:
        ancho = slide.shapes[-1].left + slide.shapes[-1].width if slide.shapes else prs.slide_width
        # Zona vacía de la plantilla: entre 43.5% y 50.5% de la altura.
        alto = prs.slide_height
        nombre_box = slide.shapes.add_textbox(
            Emu(int(prs.slide_width * 0.08)),
            Emu(int(alto * 0.435)),
            Emu(int(prs.slide_width * 0.84)),
            Emu(int(alto * 0.07)),
        )
        tf = nombre_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = nombre_titulado
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor(0x1F, 0x3D, 0x64)  # azul oscuro del diseño

    archivo = BytesIO()
    prs.save(archivo)
    archivo.seek(0)
    return archivo


def generar_certificado_pdf(inscripcion):
    """
    Genera el certificado como PDF en memoria, usando la misma imagen de fondo
    de la plantilla de PowerPoint y colocando el nombre del participante en la
    zona vacía central. Devuelve un BytesIO listo para descargarse o adjuntarse.
    Lanza ValueError si no se puede preparar el PDF.
    """
    usuario = inscripcion.usuario
    nombre_titulado = _construir_nombre_titulado(usuario)

    ruta_bg = os.path.join(settings.BASE_DIR, 'static', 'templates', '_bg_cert.png')
    if not os.path.exists(ruta_bg):
        # Fallback: extraer la imagen de fondo desde la plantilla pptx.
        ruta_plantilla = os.path.join(settings.BASE_DIR, 'static', 'templates', 'plantilla_certificado.pptx')
        if not os.path.exists(ruta_plantilla):
            raise ValueError(f"No se encontró la plantilla en {ruta_plantilla}")
        prs = Presentation(ruta_plantilla)
        extraida = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13:  # Picture
                    shape.image.save(ruta_bg)
                    extraida = True
                    break
            if extraida:
                break
        if not extraida:
            raise ValueError("La plantilla no contiene una imagen de fondo extraíble")

    try:
        bg = PILImage.open(ruta_bg)
        ancho_pag, alto_pag = bg.size
        lector = ImageReader(ruta_bg)
    except Exception as e:
        raise ValueError(f"No se pudo cargar la imagen de fondo del certificado: {e}")

    archivo = BytesIO()
    pdf = canvas.Canvas(archivo, pagesize=(ancho_pag, alto_pag))

    # Imagen de fondo a página completa
    pdf.drawImage(lector, 0, 0, width=ancho_pag, height=alto_pag, mask='auto')

    # Nombre del participante en la zona vacía central (43.5% - 50.5% de la altura)
    tamano_fuente = max(30, int(alto_pag * 0.055))
    y_nombre = int(alto_pag * 0.495)
    pdf.setFillColorRGB(0x1F / 255, 0x3D / 255, 0x64 / 255)  # azul oscuro del diseño
    pdf.setFont('Helvetica-Bold', tamano_fuente)
    pdf.drawCentredString(ancho_pag / 2, y_nombre, nombre_titulado)

    pdf.setTitle(f'Certificado de participación - {inscripcion.jornada.nombre}')
    pdf.save()
    archivo.seek(0)
    return archivo


def nombre_archivo_certificado(inscripcion):
    usuario = inscripcion.usuario
    return f"Certificado_{usuario.first_name}_{usuario.last_name}.pdf"
