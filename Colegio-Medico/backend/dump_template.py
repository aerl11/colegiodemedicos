from pptx import Presentation
import glob

p = glob.glob(r'C:\Users\PC1\Desktop\Colegio-Medico\**\*certificado participantes*.pptx', recursive=True)
print('TEMPLATE FILES:', p)
prs = Presentation(p[0])
print('slide size:', prs.slide_width, prs.slide_height, prs.slide_width/914400, 'x', prs.slide_height/914400, 'in')
for i, s in enumerate(prs.slides):
    print('SLIDE', i)
    for sh in s.shapes:
        print('  SHAPE', sh.shape_id, sh.shape_type, sh.name,
              'pos=(%.3f,%.3f)' % (sh.left/914400, sh.top/914400),
              'size=(%.3f,%.3f)' % (sh.width/914400, sh.height/914400),
              'has_tf=', getattr(sh, 'has_text_frame', False))
        if getattr(sh, 'has_text_frame', False):
            print('     TEXT:', repr(sh.text_frame.text)[:300])

# ---- also dump old in-use template with runs detail ----
old = Presentation(r'C:\Users\PC1\Desktop\Colegio-Medico\Colegio-Medico\backend\static\templates\plantilla_certificado.pptx')
print('OLD slide size:', old.slide_width/914400, 'x', old.slide_height/914400, 'in')
for sh in old.slides[0].shapes:
    print('  OLD SHAPE', sh.shape_id, sh.shape_type, sh.name,
          'pos=(%.3f,%.3f)' % (sh.left/914400, sh.top/914400),
          'size=(%.3f,%.3f)' % (sh.width/914400, sh.height/914400))
    if getattr(sh, 'has_text_frame', False):
        for pi, para in enumerate(sh.text_frame.paragraphs):
            runs = [(r.text, r.font.size.pt if r.font.size else None, r.font.name, r.font.bold) for r in para.runs]
            print('     P%d align=%s runs=%s' % (pi, para.alignment, runs))
